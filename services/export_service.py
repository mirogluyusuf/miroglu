from pathlib import Path
import pandas as pd
from docx import Document

from database.db_manager import DatabaseManager
from services.analysis_service import AnalysisService


class ExportService:
    def __init__(self, db: DatabaseManager, analysis: AnalysisService) -> None:
        self.db = db
        self.analysis = analysis

    def export_company_csv(self, company_id: int, path: str) -> str:
        a = self.analysis.analyze_company(company_id)
        rows = [{"metric": k, "value": v} for k, v in a.get("ratios", {}).items()]
        pd.DataFrame(rows).to_csv(path, index=False)
        return path

    def export_company_excel(self, company_id: int, path: str) -> str:
        a = self.analysis.analyze_company(company_id)
        rows = [{"metric": k, "value": v} for k, v in a.get("ratios", {}).items()]
        pd.DataFrame(rows).to_excel(path, index=False)
        return path

    def export_summary_docx(self, company_id: int, path: str) -> str:
        a = self.analysis.analyze_company(company_id)
        doc = Document()
        doc.add_heading("Şirket Yatırımcı Özeti", level=1)
        snap = a.get("snapshot", {})
        doc.add_paragraph(f"Şirket: {snap.get('name', '-')}")
        doc.add_paragraph(f"Sektör: {snap.get('sector_name', '-')}")
        doc.add_paragraph(f"Skor: {a.get('score', '-')}")
        doc.add_paragraph(f"Veri Kalitesi: {a.get('quality', '-')}")
        doc.add_heading("Genel Özet", level=2)
        doc.add_paragraph(a.get("commentary", {}).get("genel_ozet", "-"))
        doc.add_heading("Güçlü Yönler", level=2)
        for line in a.get("commentary", {}).get("guclu_yonler", []):
            doc.add_paragraph(line, style="List Bullet")
        doc.add_heading("Riskler", level=2)
        for line in a.get("commentary", {}).get("riskler", []):
            doc.add_paragraph(line, style="List Bullet")
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        return path
